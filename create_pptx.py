from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import os

# Create presentation
prs = Presentation()

# Brand color (Crimson Red)
BRAND_RED = RGBColor(192, 64, 72)
BRAND_DARK = RGBColor(26, 26, 26)

logo_path = 'images/photo-1791461391-1.jpg'
has_logo = os.path.exists(logo_path)

# --- Slide 1: Portada ---
slide_layout = prs.slide_layouts[0] # Title slide
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]

title.text = "Estrategia e Integración Tecnológica: Odoo POS + Profit Plus 2k12"
title.text_frame.paragraphs[0].font.color.rgb = BRAND_RED
title.text_frame.paragraphs[0].font.bold = True

subtitle.text = "Una inversión sostenible, escalable y protegida a largo plazo.\n\nPresentado a: Junta Directiva y Dueños de la Empresa."

if has_logo:
    slide.shapes.add_picture(logo_path, Inches(4), Inches(0.5), width=Inches(2))

def add_standard_slide(title_text, content_list):
    layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(layout)
    title = slide.shapes.title
    title.text = title_text
    title.text_frame.paragraphs[0].font.color.rgb = BRAND_RED
    
    tf = slide.placeholders[1].text_frame
    for i, content in enumerate(content_list):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = content
        p.font.size = Pt(20)
        p.level = 0
        if ":" in content:
            # simple bolding simulation by splitting, though python-pptx needs runs.
            pass

    if has_logo:
        slide.shapes.add_picture(logo_path, Inches(8.5), Inches(0.2), width=Inches(1.2))
    return slide, tf

def add_bold_run(paragraph, text, is_bold=True):
    run = paragraph.add_run()
    run.text = text
    run.font.bold = is_bold
    return run

# --- Slide 2 ---
layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(layout)
title = slide.shapes.title
title.text = "La Estrategia General (El Modelo Híbrido)"
title.text_frame.paragraphs[0].font.color.rgb = BRAND_RED
tf = slide.placeholders[1].text_frame
tf.clear() # clear empty p

p = tf.paragraphs[0]
add_bold_run(p, "Profit Plus 2k12 (Core Administrativo/Fiscal): ")
p.add_run().text = "El motor central probado, robusto y adaptado a las normativas legales y contables."

p = tf.add_paragraph()
add_bold_run(p, "Odoo POS (Punto de Venta Dinámico): ")
p.add_run().text = "La interfaz moderna, rápida e intuitiva para la atención y cobro al cliente."

p = tf.add_paragraph()
add_bold_run(p, "Protección de la Inversión: ")
p.add_run().text = "Combinación de plataformas consolidadas con abundante personal técnico y operativo en el mercado laboral (sin dependencia de proveedores únicos)."

if has_logo: slide.shapes.add_picture(logo_path, Inches(8.5), Inches(0.2), width=Inches(1.2))

# --- Slide 3 ---
slide, tf = add_standard_slide("Eficiencia Financiera y Licenciamiento", [])
tf.clear()
p = tf.paragraphs[0]
add_bold_run(p, "Licencias Ilimitadas: ")
p.add_run().text = "Esquema multiempresa y multiusuario sin restricción de número de cajeros o administrativos."
p = tf.add_paragraph()
add_bold_run(p, "Crecimiento a Bajo Costo: ")
p.add_run().text = "Apertura de nuevas sucursales o empresas del grupo sin pagar licencias adicionales por usuario."
p = tf.add_paragraph()
add_bold_run(p, "Retorno de Inversión (ROI): ")
p.add_run().text = "Dilución del costo fijo a medida que aumenta el volumen de ventas y operaciones."

# --- Slide 4 ---
slide, tf = add_standard_slide("Claves Operativas y Control Comercial", [])
tf.clear()
p = tf.paragraphs[0]
add_bold_run(p, "Continuidad Operativa (Modo Offline): ")
p.add_run().text = "Si se cae el internet o la red local, las cajas siguen vendiendo y se sincronizan al restablecer la conexión."
p = tf.add_paragraph()
add_bold_run(p, "Realidad Nacional y Pagos: ")
p.add_run().text = "Integración con puntos de venta bancarios, Cashea, tasa de cambio dinámica por documento y registro exacto de vueltos por Pago Móvil."
p = tf.add_paragraph()
add_bold_run(p, "Gestión de Inventario Especializada:")
p = tf.add_paragraph()
p.level = 1
add_bold_run(p, "Restaurantes: ", False).font.italic = True
p.add_run().text = "Descargo automático de recetas e ingredientes."
p = tf.add_paragraph()
p.level = 1
add_bold_run(p, "Farmacias: ", False).font.italic = True
p.add_run().text = "Control de lotes, vencimientos, fraccionamiento (caja/blíster) y principios activos."

# --- Slide 5 ---
slide, tf = add_standard_slide("Control, Auditoría y Prevención de Fraude", [])
tf.clear()
p = tf.paragraphs[0]
add_bold_run(p, "Cero Texto Libre: ")
p.add_run().text = "Clasificación obligatoria mediante listas desplegables en caja chica, gastos y compras."
p = tf.add_paragraph()
add_bold_run(p, "Trazabilidad 100%: ")
p.add_run().text = "Mapeo de cajeros reales (sin usuarios genéricos) y código único por transacción (Correlation ID)."
p = tf.add_paragraph()
add_bold_run(p, "Integración Segura (Middleware): ")
p.add_run().text = "Comunicación cifrada sin alteraciones manuales en las bases de datos."

# --- Slide 6 ---
slide, tf = add_standard_slide("Modelo de Gobierno y Estructura de TI", [])
tf.clear()
p = tf.paragraphs[0]
add_bold_run(p, "Director de Sistemas (Estrategia y Orquestación): ")
p.add_run().text = "Supervisa el puente de integración, la seguridad de la información y la evolución del sistema."
p = tf.add_paragraph()
add_bold_run(p, "Asistentes In-Situ (Soporte Directo): ")
p.add_run().text = "Atención inmediata en sucursales para asegurar que el punto de venta nunca se detenga."

# --- Slide 7 ---
slide, tf = add_standard_slide("Conclusión", [])
tf.clear()
p = tf.paragraphs[0]
add_bold_run(p, "Inversión Definitiva: ")
p.add_run().text = "Solución sólida que combina la mejor tecnología del mercado con bajo riesgo operativo."
p = tf.add_paragraph()
add_bold_run(p, "Control Total: ")
p.add_run().text = "Garantía de resguardo patrimonial, trazabilidad de inventarios y flujo de caja transparente."

# Save PPTX
prs.save('Presentacion_Junta_Directiva.pptx')
print("PPTX saved successfully.")
